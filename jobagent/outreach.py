"""Local preparation state for per-job greetings and tailored resumes."""

from __future__ import annotations

import hashlib
import json
import re
import shutil
import sqlite3
import subprocess
import tempfile
from dataclasses import dataclass
from datetime import datetime, timezone
from pathlib import Path

from docx import Document
from docx.enum.text import WD_LINE_SPACING
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Pt, RGBColor
from PIL import Image, ImageChops

from jobagent.config import ROOT
from jobagent.job_access import (
    classify_send_state,
    classify_send_states,
    compute_capabilities,
    is_excluded_company,
)


CONFIRMATION_KINDS = {"greeting", "resume"}

# 内容模式：界面必须显示来源，且不能把通用模板称为“AI 定制”。
CONTENT_SOURCE_LABELS = {
    "manual": "人工精修",
    "model": "模型生成",
    "template": "已认可通用模板",
}

# 已认可通用模板招呼语：只做自我定位与低压力下一步，不编造任何经历或数字。
TEMPLATE_GREETING = (
    "你好，我目前在做 AI 应用与智能体方向的开发，看到这个岗位想进一步了解。"
    "方便的话，想请你介绍一下团队现在主要推进的事情。"
)


def _now() -> str:
    return datetime.now().astimezone().isoformat(timespec="seconds")


class OutreachValidationError(ValueError):
    pass


@dataclass(frozen=True)
class RenderedResume:
    docx_path: Path
    image_path: Path
    docx_hash: str
    image_hash: str
    image_width: int
    image_height: int


def _file_hash(path: Path) -> str:
    digest = hashlib.sha256()
    with path.open("rb") as handle:
        for chunk in iter(lambda: handle.read(1024 * 1024), b""):
            digest.update(chunk)
    return digest.hexdigest()


def normalize_resume_assets(raw, *, root: Path | None = None) -> list[dict]:
    """把 config.resume.assets 归一化成可用的简历资产列表。

    每项形如::

        - id: "fde-solution"
          label: "FDE · AI 解决方案交付"
          match_keywords: ["FDE", "解决方案", "交付"]
          docx: "output/resume-pack/01-FDE.docx"
          image: "output/resume-pack/01-FDE.jpg"

    相对路径基于项目根目录解析；文件不存在直接报错（宁可启动失败，不可发送时才发现）。
    """
    if not raw:
        return []
    if not isinstance(raw, list):
        raise OutreachValidationError("resume.assets 必须是列表")
    base = (root or ROOT).resolve()
    assets: list[dict] = []
    for index, item in enumerate(raw):
        if not isinstance(item, dict):
            raise OutreachValidationError(f"resume.assets 第 {index + 1} 项格式不正确")
        asset_id = str(item.get("id") or "").strip()
        if not asset_id:
            raise OutreachValidationError(f"resume.assets 第 {index + 1} 项缺少 id")
        docx = Path(str(item.get("docx") or ""))
        image = Path(str(item.get("image") or ""))
        if not docx.is_absolute():
            docx = base / docx
        if not image.is_absolute():
            image = base / image
        docx = docx.resolve()
        image = image.resolve()
        if not docx.is_file():
            raise OutreachValidationError(f"简历资产 {asset_id} 的 Word 文件不存在：{docx}")
        if not image.is_file():
            raise OutreachValidationError(f"简历资产 {asset_id} 的长图不存在：{image}")
        assets.append(
            {
                "id": asset_id,
                "label": str(item.get("label") or asset_id),
                "keywords": [str(k) for k in (item.get("match_keywords") or []) if str(k).strip()],
                "docx_path": docx,
                "image_path": image,
            }
        )
    ids = [asset["id"] for asset in assets]
    if len(set(ids)) != len(ids):
        raise OutreachValidationError("resume.assets 的 id 重复")
    return assets


class TailoredResumeRenderer:
    """Create a job-specific DOCX and a mobile-readable long JPEG."""

    SOFFICE_PATHS = (
        Path(r"C:\Program Files\LibreOffice\program\soffice.exe"),
        Path(r"C:\Program Files (x86)\LibreOffice\program\soffice.exe"),
    )
    PDFTOPPM_PATHS = (
        Path(
            "C:/Users/Administrator/.cache/codex-runtimes/"
            "codex-primary-runtime/dependencies/native/poppler/"
            "Library/bin/pdftoppm.exe"
        ),
    )

    def __init__(
        self,
        template_path: str | Path,
        output_dir: str | Path = "output/tailored",
        *,
        image_width: int = 1400,
    ):
        template = Path(template_path)
        if not template.is_absolute():
            template = ROOT / template
        output = Path(output_dir)
        if not output.is_absolute():
            output = ROOT / output
        if not template.is_file():
            raise OutreachValidationError(f"Word 母版不存在：{template}")
        self.template_path = template.resolve()
        self.output_dir = output.resolve()
        self.image_width = max(1000, min(int(image_width), 2000))

    @staticmethod
    def _safe_name(job: dict, artifact_key: str) -> str:
        label = f"{job.get('company', '')}-{job.get('title', '')}-{artifact_key}"
        safe = re.sub(r"[^\w\u4e00-\u9fff-]+", "_", label).strip("_")
        return safe[:120] or artifact_key

    @staticmethod
    def _set_font(run, *, size: float, bold: bool = False, color: str = "202B35") -> None:
        run.bold = bold
        run.font.name = "Microsoft YaHei"
        run._element.get_or_add_rPr().rFonts.set(qn("w:eastAsia"), "Microsoft YaHei")
        run.font.size = Pt(size)
        run.font.color.rgb = RGBColor.from_string(color)

    @staticmethod
    def _add_left_rule(paragraph) -> None:
        properties = paragraph._p.get_or_add_pPr()
        borders = properties.find(qn("w:pBdr"))
        if borders is None:
            borders = OxmlElement("w:pBdr")
            properties.append(borders)
        left = OxmlElement("w:left")
        left.set(qn("w:val"), "single")
        left.set(qn("w:sz"), "16")
        left.set(qn("w:space"), "8")
        left.set(qn("w:color"), "13846B")
        borders.append(left)

    @staticmethod
    def _plain_markdown(value: str) -> str:
        return re.sub(r"\*\*(.+?)\*\*", r"\1", value).strip()

    def _build_docx(self, job: dict, markdown: str, destination: Path) -> None:
        document = Document(self.template_path)
        body = document._element.body
        for paragraph in list(document.paragraphs):
            body.remove(paragraph._element)

        title = next(
            (line[2:].strip() for line in markdown.splitlines() if line.startswith("# ")),
            str(job.get("title") or "AI Agent 应用工程师"),
        )
        if document.tables:
            header_cell = document.tables[0].cell(0, 0)
            if len(header_cell.paragraphs) >= 2:
                subtitle = header_cell.paragraphs[1]
                subtitle.text = title.split("|")[-1].strip() or title
                for run in subtitle.runs:
                    self._set_font(run, size=13, color="687B91")

        for raw_line in markdown.splitlines():
            line = raw_line.strip()
            if not line or line.startswith("# ") or line.startswith("<!--"):
                continue
            if line.startswith("## "):
                paragraph = document.add_paragraph()
                paragraph.paragraph_format.space_before = Pt(9)
                paragraph.paragraph_format.space_after = Pt(4)
                self._add_left_rule(paragraph)
                run = paragraph.add_run(self._plain_markdown(line[3:]))
                self._set_font(run, size=14, bold=True)
                continue
            if line.startswith("### "):
                paragraph = document.add_paragraph()
                paragraph.paragraph_format.space_before = Pt(4)
                paragraph.paragraph_format.space_after = Pt(1)
                run = paragraph.add_run(self._plain_markdown(line[4:]))
                self._set_font(run, size=10.5, bold=True)
                continue
            if line.startswith("- "):
                paragraph = document.add_paragraph(style="Resume Bullet")
                paragraph.paragraph_format.space_after = Pt(1.5)
                paragraph.paragraph_format.line_spacing_rule = WD_LINE_SPACING.SINGLE
                run = paragraph.add_run(self._plain_markdown(line[2:]))
                self._set_font(run, size=9.5)
                continue
            paragraph = document.add_paragraph()
            paragraph.paragraph_format.space_after = Pt(2)
            text = self._plain_markdown(line)
            run = paragraph.add_run(text)
            is_tech = text.startswith(("技术栈：", "技术栈:"))
            self._set_font(run, size=9 if is_tech else 9.5, color="687B91" if is_tech else "202B35")

        document.core_properties.title = (
            f"{job.get('company', '')} - {job.get('title', '')} - 定制简历"
        )
        document.save(destination)

    @staticmethod
    def _find_executable(name: str, candidates: tuple[Path, ...]) -> Path:
        found = shutil.which(name)
        if found:
            return Path(found)
        for candidate in candidates:
            if candidate.is_file():
                return candidate
        raise OutreachValidationError(f"缺少文档渲染依赖：{name}")

    @staticmethod
    def _run(command: list[str]) -> None:
        kwargs = {"capture_output": True, "text": True, "timeout": 90}
        if hasattr(subprocess, "CREATE_NO_WINDOW"):
            kwargs["creationflags"] = subprocess.CREATE_NO_WINDOW
        completed = subprocess.run(command, **kwargs)
        if completed.returncode != 0:
            detail = (completed.stderr or completed.stdout or "未知错误").strip()
            raise OutreachValidationError(f"简历渲染失败：{detail[:300]}")

    @staticmethod
    def _crop_page(image: Image.Image) -> Image.Image:
        rgb = image.convert("RGB")
        background = Image.new("RGB", rgb.size, "white")
        difference = ImageChops.difference(rgb, background).convert("L")
        bbox = difference.point(lambda value: 255 if value > 18 else 0).getbbox()
        if not bbox:
            return rgb.crop((0, 0, rgb.width, min(240, rgb.height)))
        bottom = min(rgb.height, bbox[3] + 70)
        return rgb.crop((0, 0, rgb.width, bottom))

    def _render_image(self, docx_path: Path, destination: Path) -> tuple[int, int]:
        soffice = self._find_executable("soffice", self.SOFFICE_PATHS)
        pdftoppm = self._find_executable("pdftoppm", self.PDFTOPPM_PATHS)
        with tempfile.TemporaryDirectory(prefix="jobagent-resume-") as temp_dir:
            temp = Path(temp_dir)
            self._run(
                [str(soffice), "--headless", "--convert-to", "pdf", "--outdir", str(temp), str(docx_path)]
            )
            pdf_path = temp / f"{docx_path.stem}.pdf"
            if not pdf_path.is_file():
                raise OutreachValidationError("Word 已生成，但没有得到 PDF 渲染结果")
            prefix = temp / "page"
            self._run([str(pdftoppm), "-png", "-r", "168", str(pdf_path), str(prefix)])
            page_paths = sorted(temp.glob("page-*.png"))
            if not page_paths:
                raise OutreachValidationError("PDF 已生成，但没有得到图片页面")
            if len(page_paths) > 2:
                raise OutreachValidationError("定制简历超过 2 页，请压缩内容后再生成长图")
            pages: list[Image.Image] = []
            try:
                for page_path in page_paths:
                    with Image.open(page_path) as source:
                        cropped = self._crop_page(source)
                        height = round(cropped.height * self.image_width / cropped.width)
                        pages.append(cropped.resize((self.image_width, height), Image.Resampling.LANCZOS))
                separator = 18 if len(pages) > 1 else 0
                total_height = sum(page.height for page in pages) + separator * (len(pages) - 1)
                long_image = Image.new("RGB", (self.image_width, total_height), "white")
                offset = 0
                for page in pages:
                    long_image.paste(page, (0, offset))
                    offset += page.height + separator
                long_image.save(destination, format="JPEG", quality=92, optimize=True, progressive=True)
                return long_image.size
            finally:
                for page in pages:
                    page.close()

    def render(self, job: dict, markdown: str, artifact_key: str) -> RenderedResume:
        content = str(markdown or "").strip()
        if not content:
            raise OutreachValidationError("定制简历不能为空")
        self.output_dir.mkdir(parents=True, exist_ok=True)
        safe_name = self._safe_name(job, artifact_key)
        docx_path = self.output_dir / f"{safe_name}.docx"
        image_path = self.output_dir / f"{safe_name}.jpg"
        self._build_docx(job, content, docx_path)
        width, height = self._render_image(docx_path, image_path)
        return RenderedResume(
            docx_path=docx_path,
            image_path=image_path,
            docx_hash=_file_hash(docx_path),
            image_hash=_file_hash(image_path),
            image_width=width,
            image_height=height,
        )


def _content_hash(value: str) -> str:
    normalized = re.sub(r"\s+", " ", str(value or "")).strip()
    return hashlib.sha256(normalized.encode("utf-8")).hexdigest()


def build_resume_catalog(markdown: str) -> dict:
    """Expose stable IDs for existing projects and bullets without editable text."""
    projects: list[dict] = []
    bullets: list[dict] = []
    section = ""
    entry = ""
    for index, raw_line in enumerate(str(markdown or "").splitlines()):
        line = raw_line.strip()
        if line.startswith("## "):
            section = line[3:].strip()
            entry = ""
        elif line.startswith("### "):
            entry = line[4:].strip()
            if section == "项目经历":
                projects.append(
                    {"id": f"project-{len(projects) + 1}", "heading": entry}
                )
        elif line.startswith("- "):
            bullets.append(
                {
                    "id": f"bullet-{len(bullets) + 1}",
                    "text": line[2:].strip(),
                    "section": section,
                    "entry": entry,
                    "line_index": index,
                }
            )
    return {"projects": projects, "bullets": bullets}


def _validated_priority(
    plan: dict, field: str, known_ids: set[str], *, limit: int
) -> list[str]:
    values = plan.get(field, [])
    if not isinstance(values, list) or any(not isinstance(value, str) for value in values):
        raise OutreachValidationError("AI 返回的简历微调计划格式不正确")
    if len(values) != len(set(values)):
        raise OutreachValidationError("AI 返回的简历微调计划包含重复 ID")
    if len(values) > limit:
        raise OutreachValidationError(f"AI 返回的简历微调计划最多前置 {limit} 条")
    unknown = set(values) - known_ids
    if unknown:
        raise OutreachValidationError(
            f"AI 返回的简历微调计划引用了不存在的 ID：{sorted(unknown)[0]}"
        )
    return values


def _validated_explanations(
    plan: dict, selected_ids: set[str], jd_text: str
) -> dict[str, dict]:
    values = plan.get("explanations", [])
    if not isinstance(values, list) or any(not isinstance(value, dict) for value in values):
        raise OutreachValidationError("AI 返回的简历微调解释格式不正确")
    result: dict[str, dict] = {}
    normalized_jd = re.sub(r"\s+", " ", str(jd_text or "")).strip()
    for value in values:
        item_id = str(value.get("id") or "").strip()
        if item_id not in selected_ids:
            raise OutreachValidationError("AI 返回的简历微调解释引用了未选择或不存在的 ID")
        if item_id in result:
            raise OutreachValidationError("AI 返回的简历微调解释包含重复 ID")
        reason = str(value.get("reason") or "").strip()
        evidence = str(value.get("jd_evidence") or "").strip()
        if not reason or len(reason) > 160 or not evidence or len(evidence) > 240:
            raise OutreachValidationError("AI 返回的简历微调解释内容无效")
        if re.sub(r"\s+", " ", evidence) not in normalized_jd:
            raise OutreachValidationError("AI 返回的 JD 依据不在岗位原文中")
        result[item_id] = {"reason": reason, "jd_evidence": evidence}
    return result


def build_resume_change_summary(
    markdown: str,
    plan: dict,
    catalog: dict | None = None,
    *,
    jd_text: str = "",
) -> dict:
    """Describe actual moves without allowing explanation text into the resume."""
    if not isinstance(plan, dict):
        raise OutreachValidationError("AI 返回的简历微调计划格式不正确")
    catalog = catalog or build_resume_catalog(markdown)
    project_priority = _validated_priority(
        plan, "project_priority", {item["id"] for item in catalog["projects"]}, limit=3
    )
    bullet_priority = _validated_priority(
        plan, "bullet_priority", {item["id"] for item in catalog["bullets"]}, limit=8
    )
    explanations = _validated_explanations(
        plan, set(project_priority) | set(bullet_priority), jd_text
    )
    changes: list[dict] = []

    project_rank = {item_id: rank for rank, item_id in enumerate(project_priority)}
    ordered_projects = sorted(
        catalog["projects"],
        key=lambda item: (
            0 if item["id"] in project_rank else 1,
            project_rank.get(item["id"], int(item["id"].split("-")[1])),
        ),
    )
    project_positions = {item["id"]: index for index, item in enumerate(ordered_projects, 1)}
    for original_position, item in enumerate(catalog["projects"], 1):
        new_position = project_positions[item["id"]]
        if original_position == new_position or item["id"] not in project_rank:
            continue
        changes.append({
            "kind": "project_moved", "id": item["id"],
            "label": item["heading"], "context": "项目经历",
            "from": original_position, "to": new_position,
            **explanations.get(item["id"], {}),
        })

    bullet_rank = {item_id: rank for rank, item_id in enumerate(bullet_priority)}
    bullet_groups: dict[tuple[str, str], list[dict]] = {}
    for item in catalog["bullets"]:
        bullet_groups.setdefault((item["section"], item["entry"]), []).append(item)
    for (section, entry), items in bullet_groups.items():
        ordered = sorted(
            items,
            key=lambda item: (
                0 if item["id"] in bullet_rank else 1,
                bullet_rank.get(item["id"], item["line_index"]),
            ),
        )
        positions = {item["id"]: index for index, item in enumerate(ordered, 1)}
        for original_position, item in enumerate(items, 1):
            new_position = positions[item["id"]]
            if original_position == new_position or item["id"] not in bullet_rank:
                continue
            changes.append({
                "kind": "bullet_moved", "id": item["id"],
                "label": item["text"], "context": entry or section,
                "from": original_position, "to": new_position,
                **explanations.get(item["id"], {}),
            })
    return {
        "changes": changes,
        "scope": {
            "projects": len(catalog["projects"]),
            "bullets": len(catalog["bullets"]),
            "rewritten": 0,
        },
    }


def compare_resume_order(source: str, tailored: str) -> dict:
    """Recover moves for packages generated before summaries were persisted."""
    source_catalog = build_resume_catalog(source)
    tailored_catalog = build_resume_catalog(tailored)
    changes: list[dict] = []

    source_projects = [item["heading"] for item in source_catalog["projects"]]
    tailored_projects = [item["heading"] for item in tailored_catalog["projects"]]
    if sorted(source_projects) == sorted(tailored_projects):
        target_positions = {label: index for index, label in enumerate(tailored_projects, 1)}
        for original_position, label in enumerate(source_projects, 1):
            new_position = target_positions[label]
            if new_position < original_position:
                changes.append({
                    "kind": "project_moved", "label": label,
                    "context": "项目经历", "from": original_position, "to": new_position,
                })

    def grouped(catalog: dict) -> dict[tuple[str, str], list[str]]:
        result: dict[tuple[str, str], list[str]] = {}
        for item in catalog["bullets"]:
            result.setdefault((item["section"], item["entry"]), []).append(item["text"])
        return result

    source_groups = grouped(source_catalog)
    tailored_groups = grouped(tailored_catalog)
    for context, labels in source_groups.items():
        target_labels = tailored_groups.get(context, [])
        if sorted(labels) != sorted(target_labels):
            continue
        target_positions = {label: index for index, label in enumerate(target_labels, 1)}
        for original_position, label in enumerate(labels, 1):
            new_position = target_positions[label]
            if new_position < original_position:
                changes.append({
                    "kind": "bullet_moved", "label": label,
                    "context": context[1] or context[0],
                    "from": original_position, "to": new_position,
                })
    return {
        "changes": changes,
        "scope": {
            "projects": len(source_catalog["projects"]),
            "bullets": len(source_catalog["bullets"]),
            "rewritten": 0,
        },
    }


def apply_controlled_resume_plan(markdown: str, plan: dict, catalog: dict | None = None) -> str:
    """Reorder existing source lines only; AI-provided text can never enter the resume."""
    if not isinstance(plan, dict):
        raise OutreachValidationError("AI 返回的简历微调计划格式不正确")
    source = str(markdown or "")
    catalog = catalog or build_resume_catalog(source)
    project_priority = _validated_priority(
        plan,
        "project_priority",
        {item["id"] for item in catalog["projects"]},
        limit=3,
    )
    bullet_priority = _validated_priority(
        plan,
        "bullet_priority",
        {item["id"] for item in catalog["bullets"]},
        limit=8,
    )
    lines = source.splitlines(keepends=True)

    bullet_rank = {item_id: rank for rank, item_id in enumerate(bullet_priority)}
    bullet_groups: dict[tuple[str, str], list[dict]] = {}
    for item in catalog["bullets"]:
        bullet_groups.setdefault((item["section"], item["entry"]), []).append(item)
    for items in bullet_groups.values():
        ordered = sorted(
            items,
            key=lambda item: (
                0 if item["id"] in bullet_rank else 1,
                bullet_rank.get(item["id"], item["line_index"]),
            ),
        )
        slots = sorted(item["line_index"] for item in items)
        original_lines = {item["id"]: lines[item["line_index"]] for item in items}
        for slot, item in zip(slots, ordered):
            lines[slot] = original_lines[item["id"]]

    project_rank = {item_id: rank for rank, item_id in enumerate(project_priority)}
    section_start = next(
        (index for index, line in enumerate(lines) if line.strip() == "## 项目经历"),
        None,
    )
    if section_start is not None:
        section_end = next(
            (
                index
                for index in range(section_start + 1, len(lines))
                if lines[index].strip().startswith("## ")
            ),
            len(lines),
        )
        starts = [
            index
            for index in range(section_start + 1, section_end)
            if lines[index].strip().startswith("### ")
        ]
        if starts:
            prefix = lines[section_start + 1 : starts[0]]
            blocks = []
            for ordinal, start in enumerate(starts):
                end = starts[ordinal + 1] if ordinal + 1 < len(starts) else section_end
                blocks.append((f"project-{ordinal + 1}", lines[start:end]))
            blocks.sort(
                key=lambda item: (
                    0 if item[0] in project_rank else 1,
                    project_rank.get(item[0], int(item[0].split("-")[1])),
                )
            )
            lines[section_start + 1 : section_end] = prefix + [
                line for _, block in blocks for line in block
            ]

    result = "".join(lines)
    if sorted(result.splitlines()) != sorted(source.splitlines()):
        raise OutreachValidationError("受控微调改变了母版内容，已阻止生成")
    return result


class OutreachService:
    """Prepare outreach content locally; this class has no browser or send capability."""

    def __init__(
        self,
        db_path: str | Path,
        *,
        resume_path: str | Path,
        job_loader,
        candidate_loader,
        ai=None,
        ai_factory=None,
        knowledge,
        renderer: TailoredResumeRenderer,
        assets: list[dict] | None = None,
    ):
        resume = Path(resume_path)
        if not resume.is_absolute():
            resume = ROOT / resume
        if not resume.is_file():
            raise OutreachValidationError(f"简历事实源不存在：{resume}")
        self.resume_path = resume.resolve()
        self.store = OutreachStore(db_path)
        self.job_loader = job_loader
        self.candidate_loader = candidate_loader
        self.ai = ai
        self.ai_factory = ai_factory
        self.knowledge = knowledge
        self.renderer = renderer
        # 可选的成品简历资产（用户自制 Word/长图）；配置后发送这些文件而不是重排渲染产物。
        self.assets = list(assets or [])

    @property
    def resume_text(self) -> str:
        return self.resume_path.read_text(encoding="utf-8")

    @property
    def resume_hash(self) -> str:
        return _content_hash(self.resume_text)

    @staticmethod
    def _prepare_job(job: dict) -> dict:
        prepared = dict(job)
        prepared["jd"] = str(prepared.get("jd") or prepared.get("jd_text") or "")
        return prepared

    @staticmethod
    def _assert_preparable(job: dict) -> None:
        """准备投递包的资格：岗位存在且有 JD、未关闭。

        刻意不要求：属于旧审阅样本、匹配结论为 priority（“挑战投”同样可准备）、
        已标记想投。排除公司与重复发送在最终发送入口生效，不在这里一票否决。
        """
        if not str(job.get("jd") or job.get("jd_text") or "").strip():
            raise OutreachValidationError("岗位缺少 JD 正文，无法准备投递包")
        if str(job.get("evidence_state") or "") in {"closed", "archived"}:
            raise OutreachValidationError("岗位已关闭，不能准备投递包")

    def list_candidates(self) -> list[dict]:
        send_states = classify_send_states(self.store.path)
        candidates = []
        for raw in self.candidate_loader() or []:
            job = self._prepare_job(raw)
            platform = str(job.get("platform") or "boss")
            job_id = str(job.get("job_id") or "")
            saved = self.store.get(
                platform,
                job_id,
                jd_hash=_content_hash(job["jd"]),
                resume_hash=self.resume_hash,
            )
            excluded = is_excluded_company(job.get("company"))
            send_state = send_states.get(
                (platform, job_id), {"state": "none", "label": "未发送"}
            )
            capabilities = compute_capabilities(
                job, send_state=send_state, package=saved, excluded=excluded
            )
            source = str((saved or {}).get("content_source") or "")
            candidates.append(
                {
                    "platform": platform,
                    "job_id": job_id,
                    "title": str(job.get("title") or ""),
                    "company": str(job.get("company") or ""),
                    "location": str(job.get("location") or ""),
                    "company_size": str(job.get("company_size") or ""),
                    "industry": str(job.get("industry") or ""),
                    "score": int(job.get("score") or 0),
                    "recommendation": str(job.get("recommendation") or "unknown"),
                    "interest": (job.get("decision") or {}).get("interest") or job.get("interest"),
                    "salary_raw": str(job.get("salary_raw") or ""),
                    "status": saved["status"] if saved else "not_generated",
                    "ready": bool(saved and saved["ready"]) and not excluded,
                    "stale": bool(saved and saved["stale"]),
                    "content_source": source,
                    "content_source_label": CONTENT_SOURCE_LABELS.get(source, ""),
                    "note": self.store.get_note(platform, job_id),
                    **capabilities,
                }
            )
        return candidates

    def _load_preparable_job(self, platform: str, job_id: str) -> dict:
        job = self.job_loader(platform, job_id)
        if not job:
            raise OutreachValidationError("岗位不存在")
        prepared = self._prepare_job(job)
        self._assert_preparable(prepared)
        return prepared

    def get(self, platform: str, job_id: str) -> dict:
        job = self.job_loader(platform, job_id)
        if not job:
            raise OutreachValidationError("岗位不存在")
        job = self._prepare_job(job)
        saved = self.store.get(
            platform, job_id,
            jd_hash=_content_hash(job["jd"]),
            resume_hash=self.resume_hash,
        )
        return self._attach_job(saved, job)

    def _attach_job(self, saved: dict | None, job: dict) -> dict:
        base = {
            "platform": str(job.get("platform") or "boss"),
            "job_id": str(job.get("job_id") or ""),
            "title": str(job.get("title") or ""),
            "company": str(job.get("company") or ""),
            "location": str(job.get("location") or ""),
            "company_size": str(job.get("company_size") or ""),
            "industry": str(job.get("industry") or ""),
            "recommendation": str(job.get("recommendation") or "unknown"),
            "interest": (job.get("decision") or {}).get("interest") or job.get("interest"),
            "score": int(job.get("score") or 0),
            "salary_raw": str(job.get("salary_raw") or ""),
            "jd_text": str(job.get("jd") or ""),
            "hard_requirements": list(job.get("hard_requirements") or []),
            "strengths": list(job.get("strengths") or []),
            "gaps": list(job.get("gaps") or []),
            "uncertainties": list(job.get("uncertainties") or []),
            # 所有岗位共用同一份母版简历；"定制"只是按 JD 重排既有内容，不换简历。
            "resume_source": self.resume_path.name,
            "resume_tailoring": "reorder_only",
            "note": self.store.get_note(str(job.get("platform") or "boss"), str(job.get("job_id") or "")),
        }
        if self.assets:
            suggested = self._suggest_asset_id(job)
            base["resume_assets"] = [
                {
                    "id": asset["id"],
                    "label": asset["label"],
                    "suggested": asset["id"] == suggested,
                }
                for asset in self.assets
            ]
            base["suggested_asset_id"] = suggested
        result = {
            **base,
            **(saved or {"status": "not_generated", "ready": False, "stale": False}),
        }
        if saved and not result.get("change_summary") and result.get("tailored_md"):
            result["change_summary"] = compare_resume_order(
                self.resume_text, result["tailored_md"]
            )
        # 统一透传内容来源与操作资格：前端只读结果，不自行复制规则。
        source = str(result.get("content_source") or "")
        result["content_source_label"] = CONTENT_SOURCE_LABELS.get(source, "")
        send_state = classify_send_state(
            self.store.path, base["platform"], base["job_id"]
        )
        result["send_state"] = send_state["state"]
        result["send_state_label"] = send_state["label"]
        result.update(
            compute_capabilities(
                job,
                send_state=send_state,
                package=result if saved else None,
                excluded=is_excluded_company(base.get("company")),
            )
        )
        return result

    # ------------------------------------------------------------------
    # 成品简历资产（三选一）
    def _suggest_asset_id(self, job: dict) -> str:
        """按标题/JD 命中的关键词挑资产；标题命中权重更高，都不命中时取第一个。"""
        if not self.assets:
            return ""
        title = str(job.get("title") or "")
        body = str(job.get("jd") or job.get("jd_text") or "")
        best_id, best_score = self.assets[0]["id"], 0
        for asset in self.assets:
            keywords = [k for k in asset["keywords"] if k]
            score = 3 * sum(1 for k in keywords if k in title)
            score += sum(1 for k in keywords if k in body)
            if score > best_score:
                best_id, best_score = asset["id"], score
        return best_id

    def _asset_by_id(self, asset_id: str) -> dict:
        if not self.assets:
            raise OutreachValidationError("未配置 resume.assets，无法选择成品简历")
        for asset in self.assets:
            if asset["id"] == str(asset_id or "").strip():
                return asset
        known = "、".join(asset["id"] for asset in self.assets)
        raise OutreachValidationError(f"简历资产 {asset_id!r} 不存在，可选：{known}")

    def _asset_copy(self, job: dict, asset: dict) -> RenderedResume:
        """把成品资产复制成本岗位专属文件，避免源文件日后被改影响已确认内容。"""
        self.renderer.output_dir.mkdir(parents=True, exist_ok=True)
        key = f"{job['job_id']}-{asset['id']}"
        safe = TailoredResumeRenderer._safe_name(job, key)
        docx_path = self.renderer.output_dir / f"{safe}.docx"
        image_path = self.renderer.output_dir / f"{safe}.jpg"
        shutil.copyfile(asset["docx_path"], docx_path)
        shutil.copyfile(asset["image_path"], image_path)
        with Image.open(image_path) as source:
            width, height = source.size
        return RenderedResume(
            docx_path=docx_path,
            image_path=image_path,
            docx_hash=_file_hash(docx_path),
            image_hash=_file_hash(image_path),
            image_width=int(width),
            image_height=int(height),
        )

    def select_asset(self, platform: str, job_id: str, asset_id: str) -> dict:
        """切换岗位使用的成品简历，并使其长图需要重新确认。"""
        if not self.assets:
            raise OutreachValidationError("未配置 resume.assets，无法选择成品简历")
        job = self._load_preparable_job(platform, job_id)
        if self.store.get(platform, job_id) is None:
            raise OutreachValidationError("请先生成投递包")
        asset = self._asset_by_id(asset_id)
        rendered = self._asset_copy(job, asset)
        return self._attach_job(
            self.store.update_asset(
                platform, job_id,
                resume_asset_id=asset["id"],
                docx_path=str(rendered.docx_path),
                image_path=str(rendered.image_path),
                docx_hash=rendered.docx_hash,
                image_hash=rendered.image_hash,
                image_width=rendered.image_width,
                image_height=rendered.image_height,
            ),
            job,
        )

    def generate(self, platform: str, job_id: str, asset_id: str | None = None) -> dict:
        from jobagent.ai import ManualEntryMissing

        job = self._load_preparable_job(platform, job_id)
        resume_text = self.resume_text
        knowledge = self.knowledge.for_job(job) if self.knowledge else ""
        ai = self.ai or (self.ai_factory() if self.ai_factory else None)
        if ai is None:
            raise OutreachValidationError("投递准备尚未配置 AI 模型")
        asset = None
        if asset_id or self.assets:
            # 配置了成品资产：长图/Word 直接用选中的资产，不再按母版重排渲染。
            asset = self._asset_by_id(asset_id) if asset_id else self._asset_by_id(
                self._suggest_asset_id(job)
            )
        # 内容模式：人工精修 / 模型生成 / 已认可通用模板。人工 JSON 不要求覆盖
        # 全部岗位，缺条目时用结构化信号切到通用模板，而不是匹配中文报错文案。
        content_source = "model"
        try:
            greeting = ai.greeting(resume_text, job, knowledge)
        except ManualEntryMissing:
            content_source = "template"
            greeting = TEMPLATE_GREETING
        except ValueError as exc:
            raise OutreachValidationError(str(exc)) from exc

        if asset is None:
            catalog = build_resume_catalog(resume_text)
            try:
                plan = ai.resume_plan(resume_text, job, catalog)
            except ManualEntryMissing:
                plan = {}
                content_source = "template"
            except ValueError as exc:
                raise OutreachValidationError(
                    "AI 返回的简历微调计划无法解析，已阻止生成"
                ) from exc
            tailored = apply_controlled_resume_plan(resume_text, plan, catalog)
            change_summary = build_resume_change_summary(
                resume_text, plan, catalog, jd_text=job["jd"]
            )
        else:
            # 成品资产不接收任何改写；母版原文仅作内容存档展示。
            tailored = resume_text
            change_summary = {}
        if not str(greeting or "").strip() or not str(tailored or "").strip():
            raise OutreachValidationError("没有生成完整的招呼语和简历内容")
        if asset is None:
            rendered = self.renderer.render(
                job, tailored, f"{job_id}-generated-{_content_hash(tailored)[:12]}"
            )
            used_asset_id = ""
        else:
            rendered = self._asset_copy(job, asset)
            used_asset_id = asset["id"]
        saved = self.store.save_generation(
            platform=platform, job_id=job_id,
            jd_hash=_content_hash(job["jd"]), resume_hash=_content_hash(resume_text),
            greeting_text=greeting, tailored_md=tailored,
            docx_path=str(rendered.docx_path), image_path=str(rendered.image_path),
            docx_hash=rendered.docx_hash, image_hash=rendered.image_hash,
            image_width=rendered.image_width, image_height=rendered.image_height,
            change_summary=change_summary, resume_asset_id=used_asset_id,
            content_source=content_source,
        )
        return self._attach_job(saved, job)

    def update_greeting(self, platform: str, job_id: str, text: str) -> dict:
        job = self._load_preparable_job(platform, job_id)
        return self._attach_job(self.store.update_greeting(platform, job_id, text), job)

    def update_resume(self, platform: str, job_id: str, markdown: str) -> dict:
        job = self._load_preparable_job(platform, job_id)
        saved = self.store.get(platform, job_id)
        if saved is None:
            raise OutreachValidationError("请先生成投递包")
        rendered = self.renderer.render(
            job, markdown, f"{job_id}-edited-{_content_hash(markdown)[:12]}"
        )
        updated = self.store.update_resume(
            platform, job_id, tailored_md=markdown,
            docx_path=str(rendered.docx_path), image_path=str(rendered.image_path),
            docx_hash=rendered.docx_hash, image_hash=rendered.image_hash,
            image_width=rendered.image_width, image_height=rendered.image_height,
        )
        return self._attach_job(updated, job)

    def confirm(self, platform: str, job_id: str, kind: str) -> dict:
        job = self._load_preparable_job(platform, job_id)
        current = self.store.get(
            platform, job_id, jd_hash=_content_hash(job["jd"]), resume_hash=self.resume_hash
        )
        if current is None:
            raise OutreachValidationError("请先生成投递包")
        if current["stale"]:
            raise OutreachValidationError("JD 或简历事实源已变化，请重新生成投递包")
        return self._attach_job(self.store.confirm(platform, job_id, kind), job)

    def artifact(self, platform: str, job_id: str, kind: str) -> Path:
        current = self.get(platform, job_id)
        field = {"docx": "docx_path", "image": "image_path"}.get(kind)
        if field is None or not current.get(field):
            raise OutreachValidationError("投递文件不存在，请先生成投递包")
        path = Path(current[field]).resolve()
        if not path.is_file() or self.renderer.output_dir not in path.parents:
            raise OutreachValidationError("投递文件路径无效")
        return path

    def close(self) -> None:
        self.store.close()


class OutreachStore:
    def __init__(self, db_path: str | Path = "output/jobagent.db"):
        path = Path(db_path)
        if not path.is_absolute():
            path = ROOT / path
        path.parent.mkdir(parents=True, exist_ok=True)
        self.path = path.resolve()
        self.conn = sqlite3.connect(self.path)
        self.conn.row_factory = sqlite3.Row
        self._migrate()

    def _migrate(self) -> None:
        self.conn.execute(
            """
            CREATE TABLE IF NOT EXISTS job_outreach (
              platform TEXT NOT NULL,
              job_id TEXT NOT NULL,
              content_version INTEGER NOT NULL,
              jd_hash TEXT NOT NULL,
              resume_hash TEXT NOT NULL,
              greeting_text TEXT NOT NULL,
              tailored_md TEXT NOT NULL,
              docx_path TEXT NOT NULL,
              image_path TEXT NOT NULL,
              docx_hash TEXT NOT NULL,
              image_hash TEXT NOT NULL,
              image_width INTEGER NOT NULL,
              image_height INTEGER NOT NULL,
              greeting_confirmed_at TEXT,
              resume_confirmed_at TEXT,
              content_edited INTEGER NOT NULL DEFAULT 0,
              generated_at TEXT NOT NULL,
              updated_at TEXT NOT NULL,
              last_error TEXT NOT NULL DEFAULT '',
              change_summary_json TEXT NOT NULL DEFAULT '{}',
              resume_asset_id TEXT NOT NULL DEFAULT '',
              PRIMARY KEY(platform, job_id)
            )
            """
        )
        columns = {row[1] for row in self.conn.execute("PRAGMA table_info(job_outreach)")}
        if "content_edited" not in columns:
            self.conn.execute(
                "ALTER TABLE job_outreach ADD COLUMN content_edited INTEGER NOT NULL DEFAULT 0"
            )
        if "change_summary_json" not in columns:
            self.conn.execute(
                "ALTER TABLE job_outreach ADD COLUMN change_summary_json TEXT NOT NULL DEFAULT '{}'"
            )
        if "resume_asset_id" not in columns:
            self.conn.execute(
                "ALTER TABLE job_outreach ADD COLUMN resume_asset_id TEXT NOT NULL DEFAULT ''"
            )
        if "content_source" not in columns:
            self.conn.execute(
                "ALTER TABLE job_outreach ADD COLUMN content_source TEXT NOT NULL DEFAULT ''"
            )
        self.conn.execute("CREATE TABLE IF NOT EXISTS job_outreach_notes (platform TEXT NOT NULL, job_id TEXT NOT NULL, note TEXT NOT NULL DEFAULT '', updated_at TEXT NOT NULL, PRIMARY KEY(platform, job_id))")
        self.conn.commit()

    def get_note(self, platform: str, job_id: str) -> str:
        row = self.conn.execute("SELECT note FROM job_outreach_notes WHERE platform=? AND job_id=?", (platform, job_id)).fetchone()
        return str(row[0]) if row else ""

    def save_note(self, platform: str, job_id: str, note: str) -> str:
        value = note.strip()
        self.conn.execute("INSERT INTO job_outreach_notes(platform, job_id, note, updated_at) VALUES (?, ?, ?, ?) ON CONFLICT(platform, job_id) DO UPDATE SET note=excluded.note, updated_at=excluded.updated_at", (platform, job_id, value, datetime.now(timezone.utc).isoformat()))
        self.conn.commit()
        return value

    @staticmethod
    def _serialize(
        row: sqlite3.Row | None,
        *,
        jd_hash: str | None = None,
        resume_hash: str | None = None,
    ) -> dict | None:
        if row is None:
            return None
        result = dict(row)
        try:
            result["change_summary"] = json.loads(
                result.pop("change_summary_json", "{}")
            )
        except json.JSONDecodeError:
            result["change_summary"] = {}
        result["greeting_confirmed"] = bool(result["greeting_confirmed_at"])
        result["resume_confirmed"] = bool(result["resume_confirmed_at"])
        stale_reasons: list[str] = []
        if jd_hash is not None and jd_hash != result["jd_hash"]:
            stale_reasons.append("jd_changed")
        if resume_hash is not None and resume_hash != result["resume_hash"]:
            stale_reasons.append("resume_changed")
        result["stale_reasons"] = stale_reasons
        result["stale"] = bool(stale_reasons)
        result["ready"] = (
            result["greeting_confirmed"]
            and result["resume_confirmed"]
            and not result["stale"]
        )
        result["content_edited"] = bool(result["content_edited"])
        if result["ready"]:
            result["status"] = "ready"
        elif result["content_edited"]:
            result["status"] = "edited"
        elif result["greeting_confirmed"] or result["resume_confirmed"]:
            result["status"] = "partially_confirmed"
        else:
            result["status"] = "generated"
        return result

    def get(
        self,
        platform: str,
        job_id: str,
        *,
        jd_hash: str | None = None,
        resume_hash: str | None = None,
    ) -> dict | None:
        row = self.conn.execute(
            "SELECT * FROM job_outreach WHERE platform = ? AND job_id = ?",
            (platform, job_id),
        ).fetchone()
        return self._serialize(row, jd_hash=jd_hash, resume_hash=resume_hash)

    def save_generation(
        self,
        *,
        platform: str,
        job_id: str,
        jd_hash: str,
        resume_hash: str,
        greeting_text: str,
        tailored_md: str,
        docx_path: str,
        image_path: str,
        docx_hash: str,
        image_hash: str,
        image_width: int,
        image_height: int,
        change_summary: dict | None = None,
        resume_asset_id: str = "",
        content_source: str = "model",
    ) -> dict:
        existing = self.conn.execute(
            "SELECT content_version FROM job_outreach WHERE platform = ? AND job_id = ?",
            (platform, job_id),
        ).fetchone()
        version = int(existing[0]) + 1 if existing else 1
        generated_at = _now()
        with self.conn:
            self.conn.execute(
                """
                INSERT INTO job_outreach (
                  platform, job_id, content_version, jd_hash, resume_hash,
                  greeting_text, tailored_md, docx_path, image_path, docx_hash,
                  image_hash, image_width, image_height, greeting_confirmed_at,
                  resume_confirmed_at, content_edited, generated_at, updated_at, last_error,
                  change_summary_json, resume_asset_id, content_source
                ) VALUES (?,?,?,?,?,?,?,?,?,?,?,?,?,?,?,?,?,?, '', ?, ?, ?)
                ON CONFLICT(platform, job_id) DO UPDATE SET
                  content_version = excluded.content_version,
                  jd_hash = excluded.jd_hash,
                  resume_hash = excluded.resume_hash,
                  greeting_text = excluded.greeting_text,
                  tailored_md = excluded.tailored_md,
                  docx_path = excluded.docx_path,
                  image_path = excluded.image_path,
                  docx_hash = excluded.docx_hash,
                  image_hash = excluded.image_hash,
                  image_width = excluded.image_width,
                  image_height = excluded.image_height,
                  greeting_confirmed_at = NULL,
                  resume_confirmed_at = NULL,
                  content_edited = 0,
                  generated_at = excluded.generated_at,
                  updated_at = excluded.updated_at,
                  last_error = '',
                  change_summary_json = excluded.change_summary_json,
                  resume_asset_id = excluded.resume_asset_id,
                  content_source = excluded.content_source
                """,
                (
                    platform, job_id, version, jd_hash, resume_hash,
                    greeting_text.strip(), tailored_md.strip(), docx_path, image_path,
                    docx_hash, image_hash, int(image_width), int(image_height),
                    None, None, 0, generated_at, generated_at,
                    json.dumps(change_summary or {}, ensure_ascii=False),
                    str(resume_asset_id or ""),
                    str(content_source or ""),
                ),
            )
        return self.get(platform, job_id)

    def update_greeting(self, platform: str, job_id: str, greeting_text: str) -> dict:
        text = str(greeting_text or "").strip()
        if not text:
            raise OutreachValidationError("招呼语不能为空")
        updated_at = _now()
        with self.conn:
            cursor = self.conn.execute(
                """
                UPDATE job_outreach
                   SET greeting_text = ?, greeting_confirmed_at = NULL, content_edited = 1,
                       updated_at = ?, last_error = ''
                 WHERE platform = ? AND job_id = ?
                """,
                (text, updated_at, platform, job_id),
            )
        if cursor.rowcount != 1:
            raise OutreachValidationError("请先生成投递包")
        return self.get(platform, job_id)

    def update_resume(
        self,
        platform: str,
        job_id: str,
        *,
        tailored_md: str,
        docx_path: str,
        image_path: str,
        docx_hash: str,
        image_hash: str,
        image_width: int,
        image_height: int,
    ) -> dict:
        text = str(tailored_md or "").strip()
        if not text:
            raise OutreachValidationError("定制简历不能为空")
        updated_at = _now()
        with self.conn:
            cursor = self.conn.execute(
                """
                UPDATE job_outreach
                   SET tailored_md = ?, docx_path = ?, image_path = ?,
                       docx_hash = ?, image_hash = ?, image_width = ?,
                       image_height = ?, resume_confirmed_at = NULL, content_edited = 1,
                       updated_at = ?, last_error = ''
                 WHERE platform = ? AND job_id = ?
                """,
                (
                    text, docx_path, image_path, docx_hash, image_hash,
                    int(image_width), int(image_height), updated_at, platform, job_id,
                ),
            )
        if cursor.rowcount != 1:
            raise OutreachValidationError("请先生成投递包")
        return self.get(platform, job_id)

    def update_asset(
        self,
        platform: str,
        job_id: str,
        *,
        resume_asset_id: str,
        docx_path: str,
        image_path: str,
        docx_hash: str,
        image_hash: str,
        image_width: int,
        image_height: int,
    ) -> dict:
        """切换使用的简历资产：换文件、提升内容版本并清掉长图确认。

        提升版本是刻意的：发送授权按 (platform, job_id, content_version) 去重，
        版本不动会导致旧授权的 image_hash 与新文件永远对不上且无法重新授权。
        """
        updated_at = _now()
        with self.conn:
            existing = self.conn.execute(
                "SELECT content_version FROM job_outreach WHERE platform = ? AND job_id = ?",
                (platform, job_id),
            ).fetchone()
            if existing is None:
                raise OutreachValidationError("请先生成投递包")
            version = int(existing[0]) + 1
            cursor = self.conn.execute(
                """
                UPDATE job_outreach
                   SET resume_asset_id = ?, docx_path = ?, image_path = ?,
                       docx_hash = ?, image_hash = ?, image_width = ?,
                       image_height = ?, resume_confirmed_at = NULL, content_edited = 0,
                       content_version = ?, updated_at = ?, last_error = ''
                 WHERE platform = ? AND job_id = ?
                """,
                (
                    str(resume_asset_id or ""), docx_path, image_path, docx_hash,
                    image_hash, int(image_width), int(image_height), version,
                    updated_at, platform, job_id,
                ),
            )
            if cursor.rowcount != 1:
                raise OutreachValidationError("请先生成投递包")
        return self.get(platform, job_id)

    def confirm(self, platform: str, job_id: str, kind: str) -> dict:
        if kind not in CONFIRMATION_KINDS:
            raise OutreachValidationError("确认类型无效")
        field = "greeting_confirmed_at" if kind == "greeting" else "resume_confirmed_at"
        with self.conn:
            cursor = self.conn.execute(
                f"UPDATE job_outreach SET {field} = ?, updated_at = ? "
                "WHERE platform = ? AND job_id = ?",
                (_now(), _now(), platform, job_id),
            )
        if cursor.rowcount != 1:
            raise OutreachValidationError("请先生成投递包")
        return self.get(platform, job_id)

    def close(self) -> None:
        self.conn.close()
